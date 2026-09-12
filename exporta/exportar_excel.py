# -*- coding: utf-8 -*-
"""
Exporta o banco local de decisões (já triadas) para uma planilha Excel.

Uso:
    python -m exporta.exportar_excel
    python -m exporta.exportar_excel --saida exporta/saida/minha_planilha.xlsx
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from armazenamento.db import conexao

COLUNAS = [
    ("orgao", "Órgão"),
    ("tipo", "Tipo"),
    ("titulo", "Título"),
    ("numero", "Número"),
    ("ano", "Ano"),
    ("colegiado", "Colegiado"),
    ("relator", "Relator"),
    ("data_sessao", "Data da sessão"),
    ("temas", "Temas de finanças públicas"),
    ("assunto", "Assunto"),
    ("ementa", "Ementa / sumário"),
    ("url", "Link"),
]

PADRAO_SAIDA = Path(__file__).resolve().parent / "saida" / "jurisprudencia_financas_publicas.xlsx"


def exportar(caminho_saida: Path = PADRAO_SAIDA) -> int:
    caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Decisões"

    for col_idx, (_, titulo) in enumerate(COLUNAS, start=1):
        cel = ws.cell(row=1, column=col_idx, value=titulo)
        cel.font = Font(bold=True)
    ws.freeze_panes = "A2"

    linha = 2
    with conexao() as con:
        cur = con.execute(
            "SELECT orgao, tipo, titulo, numero, ano, colegiado, relator, "
            "data_sessao, temas_json, assunto, ementa, url FROM decisoes "
            "ORDER BY orgao, data_sessao DESC"
        )
        for row in cur:
            (orgao, tipo, titulo, numero, ano, colegiado, relator,
             data_sessao, temas_json, assunto, ementa, url) = row
            temas = ", ".join(json.loads(temas_json)) if temas_json else ""
            valores = [orgao, tipo, titulo, numero, ano, colegiado, relator,
                       data_sessao, temas, assunto, ementa, url]
            for col_idx, valor in enumerate(valores, start=1):
                ws.cell(row=linha, column=col_idx, value=valor)
            linha += 1

    larguras = [10, 18, 30, 10, 8, 16, 20, 14, 30, 30, 60, 40]
    for col_idx, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = largura

    wb.save(caminho_saida)
    return linha - 2


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--saida", type=Path, default=PADRAO_SAIDA)
    args = ap.parse_args()
    n = exportar(args.saida)
    print(f"{n} decisões exportadas para {args.saida} "
          f"({dt.datetime.now():%d/%m/%Y %H:%M})")


if __name__ == "__main__":
    main()
